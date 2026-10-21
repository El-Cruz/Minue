"""Read an authorized XLSX export without editing it; write private migration data."""
import argparse, collections, datetime, hashlib, json, pathlib, re
from urllib.parse import urlparse
from zoneinfo import ZoneInfo
import openpyxl

SOURCE_ID = '1kMNdFi8zc7ppZzOTmIqsUWcu5o9EE03prqE6yeQY2PU'
REQUIRED = {
 'Reporte Entregas': ['ID','CodigoEntrega','Fecha','Tienda','Entregado por','Recibido por','Observaciones','Firma','Estado','GenerarComprobante','ComprobantePDF'],
 'Productos Entrega': ['IDProducto','IDEntrega','CodigoEntrega','Fecha','Tienda','Receta','Kg'],
 'Recetas': ['ID Receta','Nombre Receta'],
}
def value(v):
 return v.isoformat() if isinstance(v,(datetime.datetime,datetime.date)) else v

def date(v):
 if isinstance(v,datetime.datetime): d=v
 else:
  try: d=datetime.datetime.fromisoformat(str(v))
  except ValueError: d=datetime.datetime.strptime(str(v),'%d/%m/%Y %H:%M')
 if not d.tzinfo: d=d.replace(tzinfo=ZoneInfo('America/Guayaquil'))
 return d.astimezone(datetime.timezone.utc).isoformat()

def prepare(path, species_by_name=False):
 path=pathlib.Path(path); workbook=openpyxl.load_workbook(path,data_only=False)
 cached=openpyxl.load_workbook(path,data_only=True)
 sheets={}; warnings=[]
 for sheet in workbook:
  rows=[]; formulas=[]
  for row in sheet:
   if not any(c.value is not None for c in row): continue
   vals=[]
   for c in row:
    if c.data_type=='f':
     formulas.append({'cell':c.coordinate,'formula':c.value})
     if cached[sheet.title][c.coordinate].value is None: raise ValueError('Formula without cached value: '+sheet.title+'!'+c.coordinate)
    vals.append(value(cached[sheet.title][c.coordinate].value))
   rows.append({'row':row[0].row,'values':vals})
  sheets[sheet.title]={'state':sheet.sheet_state,'dimensions':[sheet.max_row,sheet.max_column],'rows':rows,'formulas':formulas}
 if set(REQUIRED)-set(sheets): raise ValueError('Required tabs are missing.')
 extra=[name for name,s in sheets.items() if name not in REQUIRED and s['rows']]
 if extra: raise ValueError('Review additional populated tabs before importing: '+', '.join(extra))
 tables={}
 for name,columns in REQUIRED.items():
  rows=sheets[name]['rows']; header=rows[0]
  indices={v:i for i,v in enumerate(header['values']) if v is not None}
  if set(columns)!=set(indices): raise ValueError('Unexpected columns in '+name)
  tables[name]=[{'source_row':r['row'],'raw':{c:r['values'][indices[c]] if indices[c]<len(r['values']) else None for c in columns}} for r in rows[1:]]
 def unique(rows,key):
  vals=[r['raw'][key] for r in rows]
  if any(not v for v in vals) or len(vals)!=len(set(vals)): raise ValueError('Missing or duplicate '+key)
 for name,key in [('Reporte Entregas','ID'),('Productos Entrega','IDProducto'),('Recetas','ID Receta')]: unique(tables[name],key)
 products=[]
 for row in tables['Recetas']:
  r=row['raw'];name=r['Nombre Receta']
  if not isinstance(name,str) or not 2<=len(name)<=150: raise ValueError('Invalid recipe name')
  products.append({**row,'legacy_id':str(r['ID Receta']),'name':name,'species':('cat' if 'gatos' in name.casefold() else 'dog') if species_by_name else None})
 recipes={p['legacy_id']:p for p in products};deliveries={};destinations={}
 for row in tables['Reporte Entregas']:
  r=row['raw'];code=r['CodigoEntrega'];match=re.fullmatch(r'([A-Z0-9]{2,8})-(\d+)',str(code))
  if not match: raise ValueError('Invalid original delivery code')
  if r['Estado'] not in ('Pendiente','Entregado'): raise ValueError('Unknown delivery state')
  if r['ComprobantePDF']:
   u=urlparse(r['ComprobantePDF'])
   if u.scheme!='https' or u.hostname!='drive.google.com' or not re.fullmatch(r'/file/d/[A-Za-z0-9_-]+/view',u.path): raise ValueError('Unexpected PDF URL')
  name=r['Tienda'];prefix,number=match.group(1),int(match.group(2))
  if not name: raise ValueError('Missing destination')
  if name in destinations and destinations[name]['code']!=prefix: raise ValueError('Destination has conflicting prefixes')
  dest=destinations.setdefault(name,{'name':name,'code':prefix,'next_number':1});dest['next_number']=max(dest['next_number'],number+1)
  deliveries[r['ID']]={**row,'legacy_id':r['ID'],'code':code,'destination':name,'date':date(r['Fecha']),'status':{'Pendiente':'pending','Entregado':'delivered'}[r['Estado']],'courier':r['Entregado por'],'recipient':r['Recibido por'],'observation':r['Observaciones'],'signature_reference':r['Firma'],'pdf_url':r['ComprobantePDF'],'items':[]}
 for row in tables['Productos Entrega']:
  r=row['raw'];parent=deliveries.get(r['IDEntrega'])
  if parent is None or r['Receta'] not in recipes: raise ValueError('Unresolved delivery or recipe relationship')
  kg=r['Kg']
  if isinstance(kg,bool) or not isinstance(kg,(int,float)) or not 0<=kg<=100000: raise ValueError('Invalid kg')
  if kg==0: warnings.append({'kind':'ZERO_KG','sheet':'Productos Entrega','row':row['source_row'],'id':r['IDProducto']})
  if r['CodigoEntrega']!=parent['code']: warnings.append({'kind':'LINE_CODE_MISMATCH','sheet':'Productos Entrega','row':row['source_row'],'id':r['IDProducto']})
  if r['Tienda']!=parent['destination']: warnings.append({'kind':'LINE_DESTINATION_MISMATCH','sheet':'Productos Entrega','row':row['source_row'],'id':r['IDProducto']})
  if date(r['Fecha'])!=parent['date']: warnings.append({'kind':'LINE_DATE_MISMATCH','sheet':'Productos Entrega','row':row['source_row'],'id':r['IDProducto']})
  parent['items'].append({**row,'legacy_id':r['IDProducto'],'recipe_id':r['Receta'],'name':recipes[r['Receta']]['name'],'quantity':kg})
 if any(not d['items'] for d in deliveries.values()): raise ValueError('Delivery without product lines; review before importing')
 if len(set(d['code'] for d in deliveries.values()))!=len(deliveries): raise ValueError('Duplicate original delivery code')
 return {'format':'minue-legacy-v1','source_key':SOURCE_ID,'source_kind':'uploaded_xlsx','sha256':hashlib.sha256(path.read_bytes()).hexdigest(),'sheets':sheets,'products':products,'destinations':list(destinations.values()),'deliveries':list(deliveries.values()),'warnings':warnings,'species_rule':'user-authorized-name-gatos-else-dog' if species_by_name else 'unclassified'}

if __name__=='__main__':
 parser=argparse.ArgumentParser();parser.add_argument('xlsx');parser.add_argument('--output',required=True);parser.add_argument('--species-by-name',action='store_true');args=parser.parse_args()
 data=prepare(args.xlsx,args.species_by_name)
 out=pathlib.Path(args.output);out.parent.mkdir(parents=True,exist_ok=True);out.write_text(json.dumps(data,ensure_ascii=False,indent=2));out.chmod(0o600)
 print(json.dumps({'deliveries':len(data['deliveries']),'lines':sum(len(d['items']) for d in data['deliveries']),'products':len(data['products']),'destinations':len(data['destinations']),'warnings':dict(collections.Counter(w['kind'] for w in data['warnings'])),'output':str(out)},ensure_ascii=False))
