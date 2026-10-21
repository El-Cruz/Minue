import importlib.util,pathlib,tempfile,unittest
import openpyxl
spec=importlib.util.spec_from_file_location('prepare',pathlib.Path(__file__).with_name('prepare.py'));module=importlib.util.module_from_spec(spec);spec.loader.exec_module(module)
class PrepareTest(unittest.TestCase):
 def workbook(self,path):
  w=openpyxl.Workbook();w.remove(w.active)
  for title,headers in module.REQUIRED.items():
   s=w.create_sheet(title)
   if title=='Recetas':s.append([]);s.append([None]+headers);s.append([None,'Receta (Gatos)','Receta (Gatos)'])
   else:s.append(headers)
  w['Reporte Entregas'].append(['old-id','HIS-001','06/05/2026 12:00','Destino','Nombre histórico','N/A','Texto original','Images/sign.png','Entregado',False,'https://drive.google.com/file/d/synthetic_id/view'])
  w['Productos Entrega'].append(['line-id','old-id','OTHER-001','06/05/2026 12:00',None,'Receta (Gatos)',0])
  w.save(path);return w
 def test_preserves_original_fields_and_resolves_by_id(self):
  with tempfile.TemporaryDirectory() as tmp:
   p=pathlib.Path(tmp)/'source.xlsx';self.workbook(p);before=p.read_bytes();data=module.prepare(p,True)
   self.assertEqual(p.read_bytes(),before);self.assertEqual(data['products'][0]['species'],'cat')
   self.assertEqual(data['deliveries'][0]['date'],'2026-05-06T17:00:00+00:00');self.assertEqual(data['deliveries'][0]['recipient'],'N/A')
   self.assertEqual(data['deliveries'][0]['items'][0]['quantity'],0)
   self.assertEqual({x['kind'] for x in data['warnings']},{'ZERO_KG','LINE_CODE_MISMATCH','LINE_DESTINATION_MISMATCH'})
 def test_rejects_orphans_and_additional_populated_tabs(self):
  with tempfile.TemporaryDirectory() as tmp:
   p=pathlib.Path(tmp)/'source.xlsx';w=self.workbook(p);w['Productos Entrega']['B2']='missing';w.save(p)
   with self.assertRaisesRegex(ValueError,'relationship'):module.prepare(p)
   w=self.workbook(p);w.create_sheet('Otra tabla').append(['Dato importante']);w.save(p)
   with self.assertRaisesRegex(ValueError,'additional populated tabs'):module.prepare(p)
 def test_rejects_negative_weights_and_untrusted_pdf_hosts(self):
  with tempfile.TemporaryDirectory() as tmp:
   p=pathlib.Path(tmp)/'source.xlsx';w=self.workbook(p);w['Productos Entrega']['G2']=-1;w.save(p)
   with self.assertRaisesRegex(ValueError,'kg'):module.prepare(p)
   w=self.workbook(p);w['Reporte Entregas']['K2']='https://other.example/file/d/test/view';w.save(p)
   with self.assertRaisesRegex(ValueError,'PDF URL'):module.prepare(p)
if __name__=='__main__':unittest.main()
